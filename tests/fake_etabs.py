"""
ETABS simulado para testes sem o programa: imita o formato de retorno do comtypes
([saídas..., ret]) e o comportamento relevante do modelo:
  - resultados por Pier só para Piers com objetos assinalados e modelo analisado;
  - alterar Pier Label com o modelo travado: recusado (strict_lock=True) ou aceito
    mantendo os resultados (strict_lock=False);
  - SetModelIsLocked(False) apaga os resultados; RunAnalysis recalcula;
  - OpenFile volta ao estado salvo do arquivo (analisado se houver .Y00 ao lado).
"""
import random
from collections import OrderedDict
from pathlib import Path


class _Obj:
    def __init__(self, **kw):
        self.__dict__.update(kw)


class FakeModel:
    """
    data: {pier: [(story, combo, location, (P, V2, V3, T, M2, M3)), ...]} na ordem da API.
    full: {pilar: [...]} esforços do pilar completo (devolvidos quando as faces viram o pilar).
    stories: nomes de cima para baixo.  labels: Pier Labels definidos no modelo.
    """

    def __init__(self, data, full, stories, combos, labels, strict_lock=True, seed=42, geometry=None):
        self.data, self.full, self.combos = data, full, list(combos)
        self.geometry = geometry or {}                 # {pier: (ângulo, largura, espessura, x, y)}
        self.stories = list(stories)
        self.elev = {s: 3.0 * (len(stories) - i) for i, s in enumerate(stories)}
        self.labels = list(labels)
        self.strict_lock = strict_lock
        self.saved_assign = OrderedDict()               # objeto -> pier, como está no arquivo
        n = 0
        for pier in data:
            for _ in range(2):
                n += 1
                self.saved_assign[("area", f"A{n}")] = pier
        self.saved_assign[("frame", "C1")] = next(iter(data))   # um pilar-barra com Pier
        for b in ("B1", "B2", "B3"):
            self.saved_assign[("frame", b)] = "None"
        self.saved_assign[("area", "LAJE1")] = "None"
        self.seed = seed
        self.filename, self.units, self.sel = "", 6, set()
        self.locked = self.analyzed = False
        self.assign = OrderedDict(self.saved_assign)
        self.runs = self.set_pier_calls = 0
        self.open_log = []

        self.PointObj = _Obj(Count=lambda: 10)
        self.File = _Obj(OpenFile=self._open)
        self.PierLabel = _Obj(GetNameList=lambda: [len(self.labels), tuple(self.labels), 0],
                              SetPier=self._new_label, GetSectionProperties=self._section)
        self.AreaObj = self._objs("area")
        self.FrameObj = self._objs("frame")
        self.RespCombo = _Obj(GetNameList=lambda: [len(self.combos), tuple(self.combos), 0],
                              GetCaseList=lambda c: [1, (0,), ("LIVE",), (1.0,), 0])
        self.LoadCases = _Obj(GetNameList=lambda: [1, ("LIVE",), 0])
        self.Analyze = _Obj(GetCaseStatus=lambda: [1, ("LIVE",), (4 if self.analyzed else 1,), 0],
                            RunAnalysis=self._run)
        self.Story = _Obj(GetNameList=lambda: [len(self.stories), tuple(self.stories), 0],
                          GetElevation=lambda s: [self.elev[s], 0])
        self.Results = _Obj(Setup=_Obj(DeselectAllCasesAndCombosForOutput=self._desel,
                                       SetComboSelectedForOutput=self._sel,
                                       GetOptionMultiValuedCombo=lambda: [1, 0]),
                            PierForce=self._pier_force)

    # ---- modelo ----
    def GetModelFilename(self, full=True):
        return self.filename or "(Untitled)"

    def GetVersion(self):
        return ["22.0.0 (simulado)", 22.0, 0]

    def GetModelIsLocked(self):
        return self.locked

    def SetModelIsLocked(self, lock):
        if not lock:
            self.locked = self.analyzed = False
        return 0

    def GetPresentUnits(self):
        return self.units

    def SetPresentUnits(self, u):
        self.units = u
        return 0

    def _open(self, path):
        p = Path(path)
        if not p.is_file():
            return 1
        self.filename = str(p)
        self.open_log.append(str(p))
        self.assign = OrderedDict(self.saved_assign)
        self.analyzed = self.locked = p.with_suffix(".Y00").exists()
        return 0

    def _run(self):
        self.runs += 1
        self.locked = self.analyzed = True
        return 0

    # ---- Pier Labels ----
    def _modifiable(self):
        return not (self.locked and self.strict_lock)

    def _new_label(self, name):
        if not self._modifiable() or name in self.labels:
            return 1
        self.labels.append(name)
        return 0

    def _objs(self, kind):
        def names():
            n = [o for k, o in self.assign if k == kind]
            return [len(n), tuple(n), 0]

        def get(name):
            return [self.assign[(kind, name)], 0]

        def set_(name, pier, item=0):
            self.set_pier_calls += 1
            if not self._modifiable() or pier not in self.labels:
                return 1
            self.assign[(kind, name)] = pier
            return 0
        return _Obj(GetNameList=names, GetPier=get, SetPier=set_)

    def _assigned(self):
        return {p for p in self.assign.values() if p != "None"}

    def _rows_of(self, pier):
        if pier in self.full and pier not in self.data:
            return self.full[pier]
        return self.data.get(pier, [])

    def _section(self, pier):
        st = list(dict.fromkeys(r[0] for r in self._rows_of(pier))) if pier in self._assigned() else []
        n = len(st)
        a, w, t, x, y = self.geometry.get(pier, (0.0, 0.0, 0.0, 0.0, 0.0))
        f = lambda v: tuple([float(v)] * n)  # noqa: E731
        k = 1.0 if self.units == 12 else 1000.0          # comprimentos em m (Ton_m_C) ou mm
        z = f(0.0)
        return [n, tuple(st), f(a), tuple([1] * n), tuple([0] * n), f(w * k), f(t * k), f(w * k), f(t * k),
                tuple(["C40"] * n), f(x * k), f(y * k), z, f(x * k), f(y * k), z, 0]

    # ---- resultados ----
    def _desel(self):
        self.sel.clear()
        return 0

    def _sel(self, combo, flag):
        self.sel.add(combo)
        return 0

    def _pier_force(self):
        assert self.units == 12, "extração fora de Ton_m_C"
        if not self.analyzed:
            return [0] + [()] * 10 + [1]
        blocks = OrderedDict()
        for pier in sorted(self._assigned()):
            for r in self._rows_of(pier):
                if r[1] in self.sel:
                    blocks.setdefault((r[0], pier), []).append((r[0], pier) + r[1:])
        keys = list(blocks)
        random.Random(self.seed).shuffle(keys)             # blocos Story/Pier fora de ordem
        out = [(s, p, c, loc) + v for k in keys for s, p, c, loc, v in blocks[k]]
        cols = list(zip(*out)) or [()] * 10
        return [len(out)] + [tuple(c) for c in cols] + [0]


class FakeApp:
    def __init__(self, sap):
        self.SapModel = sap
        self.exited = 0

    def ApplicationExit(self, save):
        self.exited += 1
        return 0


def install(module, sap):
    """Troca o connect do exportar_pier_forces pelo modelo simulado (ETABS "iniciado por nós")."""
    app = FakeApp(sap)

    def connect(edb, allow_existing=True):
        module.open_file(sap, edb)
        return module.Session(app, sap, None, True)

    module.connect = connect
    return app


def synthetic_geometry(groups, regular):
    """
    Geometria inventada (plausível) para testes e demonstração: as faces de cada pilar poligonal
    formam L, U, C ou meandro; os demais Piers são retângulos. Não é a geometria da Torre C.
    """
    import math
    geo, x0 = {}, 0.0
    dirs = [(0, -1), (1, 0), (0, 1), (1, 0), (0, -1), (1, 0), (0, 1)]
    for base, faces in groups.items():
        x, y = x0, 0.0
        for i, face in enumerate(faces):
            dx, dy = dirs[i % len(dirs)]
            length = 2.6 if dx else 3.4
            nx, ny = x + dx * length, y + dy * length
            ang = math.degrees(math.atan2(dy, dx)) % 180
            geo[face] = (ang, length, 0.25, (x + nx) / 2, (y + ny) / 2)
            x, y = nx, ny
        x0 = x + 6
    for i, pier in enumerate(regular):
        geo[pier] = (90.0 if i % 3 == 0 else 0.0, 1.2 + 0.2 * (i % 4), 0.30, 3.0 * i, -12.0)
    return geo


def load_tower_c(xlsx: Path):
    """Dados da planilha da Torre C: faces (aba discr.) e pilares inteiros (aba Pier Forces)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheets = {}
    for sh in ("Pier Forces discr.", "Pier Forces"):
        d = OrderedDict()
        for r in list(wb[sh].iter_rows(values_only=True))[3:]:
            if r and r[0]:
                d.setdefault(str(r[2]), []).append(
                    (str(r[0]), str(r[3]), str(r[5]), tuple(float(v) for v in r[6:12])))
        sheets[sh] = d
    plan1 = [str(r[0]) for r in list(wb["Plan1"].iter_rows(values_only=True))[2:]]
    wb.close()
    stories = ["TOP"] + list(dict.fromkeys(reversed(plan1)))
    return sheets["Pier Forces discr."], sheets["Pier Forces"], stories
